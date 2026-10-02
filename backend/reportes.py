# -*- coding: utf-8 -*-
"""
reportes.py — Reporte de entradas de los clientes al catálogo (solo admin)
---------------------------------------------------------------------------
GET /api/reportes/accesos?desde=AAAA-MM-DD&hasta=AAAA-MM-DD         → JSON
GET /api/reportes/accesos/excel?desde=AAAA-MM-DD&hasta=AAAA-MM-DD   → .xlsx

Sin fechas: los últimos 7 días (hoy incluido).
Los días se cuentan en hora de Bolivia (UTC-4, sin horario de verano), no
en UTC: una entrada a las 21:00 de Bolivia es del mismo día, aunque en UTC
ya sea "mañana".
"""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse

from auth import UsuarioOut, get_admin
from db import supabase, traer_todo

router = APIRouter(prefix="/api/reportes", tags=["reportes"])

BOLIVIA    = timezone(timedelta(hours=-4))
MAX_DIAS   = 93
Admin      = Annotated[UsuarioOut, Depends(get_admin)]


# ─── Fechas ──────────────────────────────────────────────────────────────────
def _fecha(texto: str | None, defecto: date, nombre: str) -> date:
    if not texto:
        return defecto
    try:
        return date.fromisoformat(texto)
    except ValueError:
        raise HTTPException(status_code=422, detail=f"La fecha «{nombre}» debe ser AAAA-MM-DD.")


def _rango(desde: str | None, hasta: str | None) -> tuple[date, date]:
    hoy = datetime.now(BOLIVIA).date()
    h = _fecha(hasta, hoy, "hasta")
    d = _fecha(desde, h - timedelta(days=6), "desde")
    if d > h:
        raise HTTPException(status_code=422, detail="La fecha «desde» no puede ser posterior a «hasta».")
    if (h - d).days + 1 > MAX_DIAS:
        raise HTTPException(status_code=422, detail=f"El rango máximo es de {MAX_DIAS} días.")
    return d, h


def _a_bolivia(valor: str) -> datetime:
    dt = datetime.fromisoformat(str(valor).replace("Z", "+00:00"))
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(BOLIVIA)


# ─── Armado del reporte ──────────────────────────────────────────────────────
def _armar(desde: date, hasta: date) -> dict:
    ini = datetime.combine(desde, time.min, tzinfo=BOLIVIA).isoformat()
    fin = datetime.combine(hasta + timedelta(days=1), time.min, tzinfo=BOLIVIA).isoformat()

    accesos = traer_todo(
        lambda: supabase.table("accesos")
        .select("id, cliente_id, creado")
        .gte("creado", ini).lt("creado", fin)
        .order("creado").order("id")
    )
    clientes = {
        c["id"]: c
        for c in traer_todo(
            lambda: supabase.table("clientes")
            .select("id, codigo_cliente, nombre, vendedor, activo")
            .order("id")
        )
    }

    dias = [(desde + timedelta(days=i)).isoformat() for i in range((hasta - desde).days + 1)]

    por_cliente: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    ultimo: dict[str, datetime] = {}
    detalle: list[dict] = []
    for a in accesos:
        c = clientes.get(a["cliente_id"])
        if c is None:
            continue
        t = _a_bolivia(a["creado"])
        por_cliente[a["cliente_id"]][t.date().isoformat()] += 1
        ultimo[a["cliente_id"]] = max(t, ultimo.get(a["cliente_id"], t))
        detalle.append({"fecha": t.date().isoformat(), "hora": t.strftime("%H:%M"), "cliente": c})

    filas = []
    for cid, cuenta in por_cliente.items():
        c = clientes[cid]
        filas.append({
            "cliente_id":     cid,
            "codigo_cliente": c.get("codigo_cliente"),
            "nombre":         c["nombre"],
            "vendedor":       c.get("vendedor"),
            "total":          sum(cuenta.values()),
            "dias_activos":   len(cuenta),
            "ultimo_acceso":  ultimo[cid].strftime("%Y-%m-%d %H:%M"),
            "por_dia":        dict(cuenta),
        })
    filas.sort(key=lambda f: (-f["total"], f["nombre"].lower()))

    dia_entradas = {d: 0 for d in dias}
    dia_clientes = {d: 0 for d in dias}
    for cuenta in por_cliente.values():
        for d, n in cuenta.items():
            dia_entradas[d] += n
            dia_clientes[d] += 1

    activos = [c for c in clientes.values() if c.get("activo")]
    sin_ingreso = sorted(
        (c for c in activos if c["id"] not in por_cliente),
        key=lambda c: c["nombre"].lower(),
    )

    return {
        "desde": desde.isoformat(),
        "hasta": hasta.isoformat(),
        "dias":  dias,
        "resumen": {
            "entradas":             len(detalle),
            "clientes_que_entraron": len(filas),
            "clientes_activos":     len(activos),
            "sin_ingreso":          len(sin_ingreso),
        },
        "por_dia": [{"fecha": d, "entradas": dia_entradas[d], "clientes": dia_clientes[d]} for d in dias],
        "por_cliente": filas,
        "sin_ingreso": [
            {"codigo_cliente": c.get("codigo_cliente"), "nombre": c["nombre"], "vendedor": c.get("vendedor")}
            for c in sin_ingreso
        ],
        "_detalle": detalle,
    }


# ─── GET /api/reportes/accesos ───────────────────────────────────────────────
@router.get("/accesos")
def reporte_accesos(_: Admin, desde: str | None = None, hasta: str | None = None):
    d, h = _rango(desde, hasta)
    datos = _armar(d, h)
    datos.pop("_detalle")
    return datos


# ─── GET /api/reportes/accesos/excel ─────────────────────────────────────────
@router.get("/accesos/excel")
def reporte_accesos_excel(_: Admin, desde: str | None = None, hasta: str | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    d, h = _rango(desde, hasta)
    datos = _armar(d, h)

    wb = Workbook()
    cabecera = PatternFill("solid", fgColor="1A1A2E")

    def hoja(ws, titulos, filas, anchos):
        ws.append(titulos)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = cabecera
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for f in filas:
            ws.append(f)
        for i, a in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = a
        ws.freeze_panes = "A2"

    # 1. Resumen por cliente: una columna por día
    ws = wb.active
    ws.title = "Por cliente"
    dias = datos["dias"]
    hoja(
        ws,
        ["Código", "Cliente", "Vendedor", "Total", "Días que entró", "Última entrada"]
        + [f"{x[8:10]}/{x[5:7]}" for x in dias],
        [
            [f["codigo_cliente"], f["nombre"], f["vendedor"], f["total"], f["dias_activos"], f["ultimo_acceso"]]
            + [f["por_dia"].get(x, 0) or None for x in dias]
            for f in datos["por_cliente"]
        ],
        [12, 42, 22, 8, 10, 17] + [6] * len(dias),
    )

    # 2. Por día
    hoja(
        wb.create_sheet("Por día"),
        ["Fecha", "Entradas", "Clientes distintos"],
        [[x["fecha"], x["entradas"], x["clientes"]] for x in datos["por_dia"]],
        [14, 11, 18],
    )

    # 3. Detalle de cada entrada
    hoja(
        wb.create_sheet("Detalle"),
        ["Fecha", "Hora", "Código", "Cliente", "Vendedor"],
        [
            [e["fecha"], e["hora"], e["cliente"].get("codigo_cliente"), e["cliente"]["nombre"], e["cliente"].get("vendedor")]
            for e in sorted(datos["_detalle"], key=lambda e: (e["fecha"], e["hora"]))
        ],
        [12, 8, 12, 42, 22],
    )

    # 4. Clientes activos que no entraron en el período
    hoja(
        wb.create_sheet("No entraron"),
        ["Código", "Cliente", "Vendedor"],
        [[c["codigo_cliente"], c["nombre"], c["vendedor"]] for c in datos["sin_ingreso"]],
        [12, 42, 22],
    )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    nombre = f"accesos_catalogo_{d.isoformat()}_a_{h.isoformat()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )
