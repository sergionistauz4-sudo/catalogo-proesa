# -*- coding: utf-8 -*-
"""
excel_datos.py — Lectura y limpieza de los Excel de clientes y productos
--------------------------------------------------------------------------
Lo usan importar_clientes.py e importar_productos.py. No toca la base.

(También se acepta .csv con las mismas columnas.)

Excel de CLIENTES (primera hoja, encabezados en la fila 1):
    Cod. Cliente | Cliente | Nit | Amb. Compra | Vendedor
Excel de PRODUCTOS (hoja "OFICIAL", o "Hoja1", o la primera; encabezados en la fila 1):
    LINEA | CLASE | CODIGO | PRODUCTO | PRECIO                       (catálogo oficial)
    LINEA | SUBCLASE | CODIGO | PRODUCTO - DESCRIPCION | STOCK | PRECIO UNIDAD   (formato anterior)
    · el orden de las filas es el orden del catálogo (campo `orden`)
    · STOCK es opcional; las columnas de imagen se ignoran (las imágenes se cargan a mano)

Los encabezados se reconocen sin importar mayúsculas, tildes ni puntos.
"""

from __future__ import annotations

from pathlib import Path

from db import limpiar_espacios, limpiar_nit, normalizar_nombre

# campo en la base → columna del Excel (normalizada: sin mayúsculas, tildes ni puntos)
COLS_CLIENTES = {
    "codigo_cliente": "cod cliente",
    "nombre":         "cliente",
    "nit":            "nit",
    "ambito":         "amb compra",
    "vendedor":       "vendedor",
}
# Cada campo puede venir con distintos nombres de columna (formato nuevo y anterior).
# `stock` es opcional.
COLS_PRODUCTOS = {
    "linea":    ("linea",),
    "subclase": ("clase", "subclase"),
    "codigo":   ("codigo",),
    "nombre":   ("producto", "producto descripcion"),
    "precio":   ("precio", "precio unidad"),
    "stock":    ("stock",),
}
COLS_PRODUCTOS_OPCIONALES = {"stock"}


# ─── Helpers ─────────────────────────────────────────────────────────────────
def _arreglar_codificacion(texto: str) -> str:
    """
    Repara texto con tildes/ñ "rotas" por una mala codificación
    (ej. 'PRACTIPAÃ‘AL' → 'PRACTIPAÑAL'). Si no hace falta, lo deja igual.
    """
    if "Ã" not in texto and "Â" not in texto:
        return texto
    try:
        return texto.encode("cp1252").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return texto


def _texto(valor) -> str:
    """Celda → texto limpio. Los números enteros no llevan '.0' (50024.0 → '50024')."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return limpiar_espacios(_arreglar_codificacion(str(valor)))


def _hojas(ruta: Path, hoja: str | None):
    if ruta.suffix.lower() == ".csv":
        import csv
        texto = ruta.read_text(encoding="utf-8-sig", errors="replace")
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
        return [fila for fila in csv.reader(texto.splitlines(), dialecto)]
    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[hoja] if hoja else wb.worksheets[0]
    return [list(fila) for fila in ws.iter_rows(values_only=True)]


def _filas_como_dicts(ruta: Path, hoja: str | None, columnas: dict, opcionales: set = frozenset()):
    """
    `columnas`: campo → nombre de columna (str) o varios nombres posibles (tupla).
    Los campos en `opcionales` pueden faltar del Excel (llegan como None).
    """
    filas = _hojas(ruta, hoja)
    if not filas:
        raise SystemExit("El archivo está vacío.")
    encabezado = [normalizar_nombre(c if c is not None else "") for c in filas[0]]
    indices: dict[str, int | None] = {}
    faltan = []
    for campo, nombres in columnas.items():
        nombres = (nombres,) if isinstance(nombres, str) else nombres
        hallado = next((encabezado.index(n) for n in nombres if n in encabezado), None)
        if hallado is None and campo not in opcionales:
            faltan.append(nombres[0])
        indices[campo] = hallado
    if faltan:
        raise SystemExit(
            f"Faltan columnas en el Excel: {faltan}\nEncabezados leídos: {filas[0]}"
        )
    for n, fila in enumerate(filas[1:], start=2):
        if not any(c not in (None, "") for c in fila):
            continue  # fila totalmente vacía
        yield n, {
            campo: (fila[i] if i is not None and i < len(fila) else None)
            for campo, i in indices.items()
        }


def convertir_stock(valor) -> int | None:
    """
    El Excel trae los miles con punto y Excel los guardó como decimales:
    4.159 es en realidad 4159 unidades, y 348 es 348. Regla: un número
    con decimales se multiplica por 1000; un entero se deja igual.
    """
    if valor in (None, ""):
        return None
    try:
        n = float(str(valor).replace(",", "."))
    except ValueError:
        return None
    return round(n * 1000) if not n.is_integer() else int(n)


# ─── Clientes ────────────────────────────────────────────────────────────────
def leer_clientes(ruta: str | Path, hoja: str | None = None):
    """Devuelve (registros, avisos). Los repetidos por código se reportan y se usa el último."""
    registros: dict[str, dict] = {}
    avisos: list[str] = []
    for n, f in _filas_como_dicts(Path(ruta), hoja, COLS_CLIENTES):
        nombre = limpiar_espacios(_texto(f["nombre"]))
        nit    = limpiar_nit(_texto(f["nit"]))
        codigo = _texto(f["codigo_cliente"]).upper() or None
        if not nombre or not nit:
            avisos.append(f"fila {n}: falta nombre o NIT → se saltea ({list(f.values())})")
            continue
        if not codigo:
            avisos.append(f"fila {n}: «{nombre}» no tiene código de cliente → se saltea")
            continue
        if codigo in registros:
            avisos.append(f"fila {n}: código {codigo} repetido → se usa la última aparición")
        registros[codigo] = {
            "codigo_cliente": codigo,
            "nombre":   nombre,
            "nit":      nit,
            "ambito":   _texto(f["ambito"]).upper() or None,
            "vendedor": _texto(f["vendedor"]).upper() or None,
        }
    return list(registros.values()), avisos


# ─── Productos ───────────────────────────────────────────────────────────────
def _hoja_productos(ruta: Path) -> str | None:
    """'OFICIAL' si existe, si no 'Hoja1', si no la primera hoja."""
    if ruta.suffix.lower() == ".csv":
        return None
    from openpyxl import load_workbook
    nombres = load_workbook(ruta, read_only=True).sheetnames
    for preferida in ("OFICIAL", "Hoja1"):
        if preferida in nombres:
            return preferida
    return None


def leer_productos(ruta: str | Path, hoja: str | None = None):
    """
    Devuelve (registros, avisos). Cada registro lleva `orden` = posición en el
    Excel (1, 2, 3…), que es el orden en que se muestra el catálogo.
    """
    ruta = Path(ruta)
    hoja = hoja or _hoja_productos(ruta)
    registros: dict[str, dict] = {}
    avisos: list[str] = []
    for n, f in _filas_como_dicts(ruta, hoja, COLS_PRODUCTOS, COLS_PRODUCTOS_OPCIONALES):
        nombre = _texto(f["nombre"])
        codigo = _texto(f["codigo"]) or None
        try:
            precio = round(float(str(f["precio"]).replace(",", ".")), 2)
        except (TypeError, ValueError):
            precio = None
        if not nombre or not codigo or precio is None or precio < 0:
            avisos.append(f"fila {n}: falta nombre, código o precio válido → se saltea ({list(f.values())})")
            continue
        if codigo in registros:
            avisos.append(f"fila {n}: código {codigo} repetido → se usa la última aparición")
            del registros[codigo]          # y pasa a la posición de esta aparición
        linea    = _texto(f["linea"]).upper() or None
        subclase = _texto(f["subclase"]) or None
        registros[codigo] = {
            "codigo":      codigo,
            "linea":       linea,
            "subclase":    subclase,
            "nombre":      nombre,
            # descripción inicial = el nombre (editable desde la app; una reimportación no la pisa)
            "descripcion": nombre,
            "precio":      precio,
            "stock":       convertir_stock(f["stock"]),   # None si el Excel no trae stock
        }
    lista = list(registros.values())
    for pos, r in enumerate(lista, start=1):
        r["orden"] = pos
    return lista, avisos


# ─── Lista de precios mensual ────────────────────────────────────────────────
COLUMNA_PRECIO_NUEVO = 6      # columna G (A=0 … G=6): "NUEVOS PRECIO KC"
COLUMNA_CODIGO_DEFECTO = 2    # columna C, si no se encuentra el encabezado "CODIGO"


def _a_precio(valor) -> tuple[float | None, str | None]:
    """Celda → (precio, motivo_si_no_sirve). Acepta 126.76, '126.76' y '126,76'."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None, "celda vacía"
    if isinstance(valor, bool):
        return None, "valor no numérico"
    if isinstance(valor, (int, float)):
        n = float(valor)
    else:
        t = str(valor).strip().replace(" ", "")
        if t.startswith("#"):
            return None, f"sin precio nuevo ({t})"
        if "," in t and "." in t:          # 1.234,50
            t = t.replace(".", "").replace(",", ".")
        else:
            t = t.replace(",", ".")
        try:
            n = float(t)
        except ValueError:
            return None, f"no es un número («{str(valor).strip()[:20]}»)"
    if n != n or n <= 0 or n > 10_000_000:
        return None, "precio no válido"
    return round(n, 2), None


def leer_precios(fuente, nombre_archivo: str = ""):
    """
    Lee la lista de precios del mes. De TODO el Excel solo interesan dos cosas
    (primera hoja, encabezados en la fila 1):
        · el CÓDIGO del producto  (columna con encabezado "CODIGO"; si no, la C)
        · el precio nuevo         (columna G)

    `fuente` es una ruta o un archivo en memoria. Devuelve:
        {"hoja", "filas", "precios": {codigo: {"precio", "fila"}},
         "ignorados": [{"codigo", "fila", "motivo"}], "avisos": [...]}
    Las filas sin precio válido (#N/A, vacías…) se informan en `ignorados` y NO
    se usan: esos productos conservan su precio actual.
    """
    from openpyxl import load_workbook
    try:
        wb = load_workbook(fuente, read_only=True, data_only=True)
    except Exception as e:   # noqa: BLE001 — archivo corrupto / no es .xlsx
        raise ValueError("No pude abrir el archivo. Tiene que ser un Excel .xlsx "
                         "(si es .xlsb, abrilo en Excel y guardalo como .xlsx).") from e
    ws = wb.worksheets[0]
    filas = [list(f) for f in ws.iter_rows(values_only=True)]
    if len(filas) < 2:
        raise ValueError("La primera hoja del Excel está vacía.")

    encabezado = [normalizar_nombre(c if c is not None else "") for c in filas[0]]
    col_codigo = encabezado.index("codigo") if "codigo" in encabezado else COLUMNA_CODIGO_DEFECTO
    if max(len(f) for f in filas) <= COLUMNA_PRECIO_NUEVO:
        raise ValueError("El Excel no llega a la columna G, donde van los precios nuevos.")

    precios: dict[str, dict] = {}
    ignorados: list[dict] = []
    avisos: list[str] = []
    for n, fila in enumerate(filas[1:], start=2):
        if not any(c not in (None, "") for c in fila):
            continue
        codigo = _texto(fila[col_codigo] if col_codigo < len(fila) else None).upper()
        valor  = fila[COLUMNA_PRECIO_NUEVO] if COLUMNA_PRECIO_NUEVO < len(fila) else None
        precio, motivo = _a_precio(valor)
        if not codigo:
            if precio is not None:
                ignorados.append({"codigo": "", "fila": n, "motivo": "fila sin código"})
            continue
        if precio is None:
            ignorados.append({"codigo": codigo, "fila": n, "motivo": motivo})
            continue
        if codigo in precios:
            avisos.append(f"fila {n}: el código {codigo} está repetido → se usa la última aparición")
        precios[codigo] = {"precio": precio, "fila": n}
    return {
        "hoja": ws.title,
        "filas": len(filas) - 1,
        "precios": precios,
        "ignorados": ignorados,
        "avisos": avisos,
    }