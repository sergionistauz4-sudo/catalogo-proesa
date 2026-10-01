# -*- coding: utf-8 -*-
"""
importar_clientes.py — Carga masiva de clientes desde un Excel o CSV
---------------------------------------------------------------------
Uso:
  python importar_clientes.py clientes.xlsx              # importa
  python importar_clientes.py clientes.csv --prueba      # solo muestra qué haría

El archivo necesita (en la primera fila) una columna de nombre y una
de NIT. Se reconocen estos encabezados (sin importar mayúsculas/tildes):
  nombre: "nombre", "razon social", "cliente", "farmacia", "nombre cliente"
  NIT:    "nit", "ci/nit", "nit/ci", "nit o ci", "ci"

Si un NIT ya existe en la base, se ACTUALIZA el nombre (no se duplica).
Filas sin nombre o sin NIT se saltean y se informan al final.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from db import limpiar_espacios, limpiar_nit, normalizar_nombre, supabase

COLS_NOMBRE = {"nombre", "razon social", "cliente", "farmacia", "nombre cliente", "razonsocial"}
COLS_NIT    = {"nit", "ci nit", "cinit", "nit ci", "nitci", "nit o ci", "ci"}


def _leer_filas(ruta: Path) -> list[list[str]]:
    if ruta.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        hoja = load_workbook(ruta, read_only=True, data_only=True).active
        return [["" if v is None else str(v) for v in fila] for fila in hoja.iter_rows(values_only=True)]

    texto = ruta.read_text(encoding="utf-8-sig", errors="replace")
    dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
    return [fila for fila in csv.reader(texto.splitlines(), dialecto)]


def _numero_a_texto(valor: str) -> str:
    """Excel a veces trae el NIT como 1234567.0 — le sacamos el .0"""
    return valor[:-2] if valor.endswith(".0") and valor[:-2].isdigit() else valor


def main() -> None:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    prueba = "--prueba" in sys.argv
    if not args:
        print(__doc__)
        sys.exit(1)

    ruta = Path(args[0])
    filas = _leer_filas(ruta)
    if not filas:
        sys.exit("El archivo está vacío.")

    encabezado = [normalizar_nombre(c) for c in filas[0]]
    try:
        i_nombre = next(i for i, c in enumerate(encabezado) if c in COLS_NOMBRE)
        i_nit    = next(i for i, c in enumerate(encabezado) if c in COLS_NIT)
    except StopIteration:
        sys.exit(f"No encontré las columnas de nombre y NIT. Encabezados leídos: {filas[0]}")

    registros: dict[str, dict] = {}
    salteadas: list[str] = []
    for n, fila in enumerate(filas[1:], start=2):
        nombre = limpiar_espacios(fila[i_nombre] if i_nombre < len(fila) else "")
        nit    = limpiar_nit(_numero_a_texto(fila[i_nit].strip() if i_nit < len(fila) else ""))
        if not nombre or not nit:
            salteadas.append(f"fila {n}: {fila}")
            continue
        if nit in registros:
            salteadas.append(f"fila {n}: NIT {nit} repetido (se usa la última aparición)")
        # Sin "activo": los nuevos quedan activos (default de la tabla) y
        # a los existentes no se les toca — no reactivamos a nadie sin querer.
        registros[nit] = {"nombre": nombre, "nit": nit}

    lote = list(registros.values())
    print(f"Clientes válidos: {len(lote)}  ·  filas salteadas: {len(salteadas)}")
    for s in salteadas:
        print("  ⚠", s)

    if prueba:
        for r in lote[:15]:
            print(f"  {r['nit']:>14}  {r['nombre']}")
        print("Modo prueba: no se guardó nada.")
        return

    for i in range(0, len(lote), 500):
        supabase.table("clientes").upsert(lote[i:i + 500], on_conflict="nit").execute()
    print(f"✓ {len(lote)} clientes importados/actualizados.")


if __name__ == "__main__":
    main()
